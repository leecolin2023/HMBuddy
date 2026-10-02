"""PPTX Adapter（规格第 14 节）。

按 slide 顺序保留 slide index / title / textbox / table / 图片引用；
不识别复杂视觉布局语义，但必须知道内容位于第几页（slide attribution）。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from workspace.artifact import Artifact, ArtifactBlock, ArtifactLocator
from workspace.errors import EncryptedArtifactError

from .base import ArtifactAdapter, assign_block_ids, ensure_not_ole


def _slide_title(slide) -> str | None:
    try:
        title_shape = slide.shapes.title
    except Exception:
        return None
    if title_shape is None or not title_shape.has_text_frame:
        return None
    text = title_shape.text.strip()
    return text or None


class PptxAdapter(ArtifactAdapter):
    artifact_type = "pptx"
    supported_extensions = (".pptx",)
    parser_library = "python-pptx"

    def read(self, path: Path, artifact_id: str) -> Artifact:
        ensure_not_ole(path, self)
        prs = Presentation(str(path))

        blocks: list[ArtifactBlock] = []
        slides_meta: list[dict] = []

        def walk_shapes(shapes, slide_index: int, meta: dict, counter: list[int]) -> None:
            """深度优先遍历（含组合形状），shape_index 反映形状在页内的出现顺序。"""
            for shape in shapes:
                counter[0] += 1
                shape_index = counter[0]
                try:
                    shape_type = shape.shape_type
                except Exception:
                    shape_type = None

                if shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk_shapes(shape.shapes, slide_index, meta, counter)
                    continue

                if shape.has_table:
                    table = shape.table
                    cells = [
                        [cell.text.strip() for cell in row.cells] for row in table.rows
                    ]
                    meta["table_count"] += 1
                    blocks.append(
                        ArtifactBlock(
                            "",
                            "table",
                            None,
                            {
                                "slide": slide_index,
                                "shape_index": shape_index,
                                "table_index": meta["table_count"],
                            },
                            {
                                "rows": len(cells),
                                "columns": len(cells[0]) if cells else 0,
                                "cells": cells,
                            },
                            locator=ArtifactLocator(
                                "pptx",
                                {"slide": slide_index, "shape_index": shape_index},
                            ),
                        )
                    )
                    continue

                if shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
                    meta["image_count"] += 1
                    blocks.append(
                        ArtifactBlock(
                            "",
                            "image_reference",
                            None,
                            {"slide": slide_index, "shape_index": shape_index},
                            {"name": shape.name, "shape_type": "PICTURE"},
                            locator=ArtifactLocator(
                                "pptx",
                                {"slide": slide_index, "shape_index": shape_index},
                            ),
                        )
                    )
                    continue

                if shape.has_text_frame:
                    text = shape.text_frame.text
                    if text.strip():
                        meta["textbox_count"] += 1
                        blocks.append(
                            ArtifactBlock(
                                "",
                                "textbox",
                                text,
                                {"slide": slide_index, "shape_index": shape_index},
                                {
                                    "name": shape.name,
                                    "shape_type": str(shape_type),
                                    "is_placeholder": shape.is_placeholder,
                                },
                                locator=ArtifactLocator(
                                    "pptx",
                                    {"slide": slide_index, "shape_index": shape_index},
                                ),
                            )
                        )

        for slide_index, slide in enumerate(prs.slides, start=1):
            title = _slide_title(slide)
            meta = {
                "index": slide_index,
                "title": title,
                "textbox_count": 0,
                "table_count": 0,
                "image_count": 0,
            }
            blocks.append(
                ArtifactBlock(
                    "",
                    "slide",
                    title,
                    {"slide": slide_index},
                    {"title": title, "layout": slide.slide_layout.name},
                    locator=ArtifactLocator("pptx", {"slide": slide_index}),
                )
            )
            walk_shapes(slide.shapes, slide_index, meta, [0])
            slides_meta.append(meta)

        metadata = {
            "title": slides_meta[0]["title"] if slides_meta else None,
            "slide_count": len(slides_meta),
            "slides": slides_meta,
        }

        content_lines: list[str] = []
        for block in blocks:
            if block.block_type == "slide":
                header = f"[Slide {block.location.get('slide')}] {block.text or ''}"
                content_lines.append(header.rstrip())
            elif block.block_type == "textbox" and block.text:
                content_lines.append(block.text)
        content = "\n".join(content_lines)

        assign_block_ids(blocks)
        return self.build_artifact(
            path,
            artifact_id,
            blocks=blocks,
            metadata=metadata,
            content=content,
            file_stat=path.stat(),
        )
