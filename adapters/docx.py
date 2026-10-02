"""DOCX Adapter（规格第 11 节）。

保留 heading（含 level）/ paragraph / 基础列表 / table（rows、columns、cells）/
图片位置引用，并保证原始阅读顺序不乱（最低要求）。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from workspace.artifact import Artifact, ArtifactBlock, ArtifactLocator
from workspace.errors import EncryptedArtifactError

from .base import ArtifactAdapter, assign_block_ids, ensure_not_ole
from .tables import ExtractedTable, TableCell, TableRow, table_to_block_metadata
from .textnorm import collapse_cell_text

_HEADING_NAME_RE = re.compile(r"^(?:heading|标题|標題)\s*(\d+)$", re.IGNORECASE)
_HEADING_ID_RE = re.compile(r"^(?:heading|标题|標題)(\d+)$", re.IGNORECASE)


def _heading_level(paragraph: Paragraph) -> int | None:
    """识别标题级别：优先样式名（Heading 1 / 标题 1），其次 outlineLvl。"""
    try:
        style = paragraph.style
    except Exception:
        style = None
    if style is not None:
        name = (style.name or "").strip()
        matched = _HEADING_NAME_RE.match(name)
        if matched:
            return int(matched.group(1))
        style_id = getattr(style, "style_id", "") or ""
        matched = _HEADING_ID_RE.match(style_id)
        if matched:
            return int(matched.group(1))
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None and outline.get(qn("w:val")) is not None:
            try:
                return int(outline.get(qn("w:val"))) + 1
            except ValueError:
                return None
    return None


def _list_level(paragraph: Paragraph) -> int | None:
    """识别基础列表：段落级 numPr，或样式名含 List（List Bullet / List Number）。"""
    try:
        style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
    except Exception:
        style_name = ""
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            ilvl = num_pr.find(qn("w:ilvl"))
            if ilvl is not None and ilvl.get(qn("w:val")) is not None:
                try:
                    return int(ilvl.get(qn("w:val")))
                except ValueError:
                    return 0
            return 0
    if "list" in style_name.lower():
        return 0
    return None


def _paragraph_images(doc, paragraph: Paragraph) -> list[str]:
    """返回该段落引用的图片在包内的路径（如 /word/media/image1.png）。"""
    try:
        rids = paragraph._p.xpath(".//a:blip/@r:embed")
    except Exception:
        return []
    names: list[str] = []
    for rid in rids or []:
        part = doc.part.related_parts.get(rid)
        if part is not None:
            names.append(str(part.partname))
    return names


def _extract_cell_text(cell) -> str:
    """单元格文本：段落 + 嵌套表格（fce 能力，python-docx 的 cell.text 不含嵌套表格）。

    嵌套表格的行展开成"单元格 | 单元格"的文本行，多行内容以" / "折叠。
    """
    from docx.oxml.ns import qn as _qn

    parts: list[str] = []
    for child in cell._tc:
        if child.tag == _qn("w:p"):
            text = collapse_cell_text(Paragraph(child, cell).text)
            if text:
                parts.append(text)
        elif child.tag == _qn("w:tbl"):
            nested_table = Table(child, cell)
            nested_rows = [
                [collapse_cell_text(inner.text) for inner in row.cells]
                for row in nested_table.rows
            ]
            parts.extend(" | ".join(row_values) for row_values in nested_rows)
    return collapse_cell_text("\n".join(parts))


def _table_to_extracted(table: Table, table_index: int) -> ExtractedTable:
    """把 python-docx 表格转成 ExtractedTable，合并单元格以 tc 元素同一性识别。"""
    tc_grid = [[cell._tc for cell in row.cells] for row in table.rows]
    row_count = len(tc_grid)
    column_count = max((len(row) for row in tc_grid), default=0)

    # 找出每个 tc 的行/列跨度
    span_info: dict[int, dict] = {}
    for row_index, row in enumerate(tc_grid):
        for column_index, tc in enumerate(row):
            key = id(tc)
            info = span_info.setdefault(
                key,
                {"min_row": row_index, "max_row": row_index, "min_col": column_index, "max_col": column_index},
            )
            info["max_row"] = max(info["max_row"], row_index)
            info["max_col"] = max(info["max_col"], column_index)

    rows: list[TableRow] = []
    emitted: set[int] = set()
    for row_index in range(row_count):
        cells: list[TableCell] = []
        seen_in_row: set[int] = set()
        for column_index, tc in enumerate(tc_grid[row_index]):
            key = id(tc)
            if key in seen_in_row or key in emitted:
                continue
            seen_in_row.add(key)
            info = span_info[key]
            anchor = table.cell(info["min_row"], info["min_col"])
            cells.append(
                TableCell(
                    row=info["min_row"] + 1,
                    column=info["min_col"] + 1,
                    text=_extract_cell_text(anchor),
                    rowspan=(info["max_row"] - info["min_row"]) + 1,
                    colspan=(info["max_col"] - info["min_col"]) + 1,
                )
            )
        cells.sort(key=lambda cell: cell.column)
        rows.append(TableRow(index=row_index + 1, cells=cells))
        emitted.update(seen_in_row)

    return ExtractedTable(
        page=0,
        table=table_index,
        rows=rows,
        column_count=column_count,
        extraction_method="python-docx",
        confidence=1.0,
    )


class DocxAdapter(ArtifactAdapter):
    artifact_type = "docx"
    supported_extensions = (".docx",)
    parser_library = "python-docx"

    def read(self, path: Path, artifact_id: str) -> Artifact:
        ensure_not_ole(path, self)
        doc = Document(str(path))

        blocks: list[ArtifactBlock] = []
        counts = {
            "heading": 0,
            "paragraph": 0,
            "list_item": 0,
            "table": 0,
            "image_reference": 0,
        }
        paragraph_index = 0
        table_index = 0
        first_heading: str | None = None

        # 直接按 body XML 子元素顺序遍历，保证 paragraph 与 table 的原始阅读顺序
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, doc)
                paragraph_index += 1
                text = paragraph.text.strip()
                images = _paragraph_images(doc, paragraph)
                if images:
                    counts["image_reference"] += 1
                    blocks.append(
                        ArtifactBlock(
                            "",
                            "image_reference",
                            None,
                            {"paragraph_index": paragraph_index},
                            {"count": len(images), "images": images},
                            locator=ArtifactLocator(
                                "docx", {"paragraph_index": paragraph_index}
                            ),
                        )
                    )
                if not text:
                    continue
                level = _heading_level(paragraph)
                if level is not None:
                    counts["heading"] += 1
                    if first_heading is None:
                        first_heading = text
                    blocks.append(
                        ArtifactBlock(
                            "",
                            "heading",
                            text,
                            {"paragraph_index": paragraph_index},
                            {"level": level},
                            locator=ArtifactLocator(
                                "docx", {"paragraph_index": paragraph_index}
                            ),
                        )
                    )
                    continue
                list_level = _list_level(paragraph)
                if list_level is not None:
                    counts["list_item"] += 1
                    blocks.append(
                        ArtifactBlock(
                            "",
                            "list_item",
                            text,
                            {"paragraph_index": paragraph_index},
                            {"list_level": list_level},
                            locator=ArtifactLocator(
                                "docx", {"paragraph_index": paragraph_index}
                            ),
                        )
                    )
                    continue
                counts["paragraph"] += 1
                blocks.append(
                    ArtifactBlock(
                        "",
                        "paragraph",
                        text,
                        {"paragraph_index": paragraph_index},
                        {},
                        locator=ArtifactLocator(
                            "docx", {"paragraph_index": paragraph_index}
                        ),
                    )
                )
            elif child.tag == qn("w:tbl"):
                table = Table(child, doc)
                table_index += 1
                extracted_table = _table_to_extracted(table, table_index)
                counts["table"] += 1
                blocks.append(
                    ArtifactBlock(
                        "",
                        "table",
                        None,
                        {"table_index": table_index},
                        table_to_block_metadata(extracted_table),
                        locator=ArtifactLocator("docx", {"table_index": table_index}),
                    )
                )

        core = doc.core_properties
        core_properties = {
            "title": core.title,
            "author": core.author,
            "subject": core.subject,
            "created": core.created.isoformat() if core.created else None,
            "modified": core.modified.isoformat() if core.modified else None,
        }
        sections = [
            {
                "index": index,
                "orientation": str(section.orientation),
                "page_width": section.page_width,
                "page_height": section.page_height,
            }
            for index, section in enumerate(doc.sections, start=1)
        ]
        metadata = {
            "title": core.title or first_heading,
            "heading_count": counts["heading"],
            "paragraph_count": counts["paragraph"],
            "list_item_count": counts["list_item"],
            "table_count": counts["table"],
            "image_reference_count": counts["image_reference"],
            "section_count": len(sections),
            "sections": sections,
            "core_properties": core_properties,
        }
        content = "\n".join(block.text for block in blocks if block.text)
        assign_block_ids(blocks)
        return self.build_artifact(
            path,
            artifact_id,
            blocks=blocks,
            metadata=metadata,
            content=content,
            file_stat=path.stat(),
        )
