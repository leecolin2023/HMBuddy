"""XLSX Adapter（规格第 13 节）。

不把 Excel 压成整表 CSV 文本，保留 Workbook → Sheet → Used Range → Cell 层级；
单元格记录 value / formula / row / column；merged range 信息单独保留。
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from workspace.artifact import Artifact, ArtifactBlock
from workspace.errors import EncryptedArtifactError

from .base import ArtifactAdapter, assign_block_ids, ensure_not_ole

# Phase 1 保护性上限：超出部分标记 truncated，未来再做分段读取（规格 ER-04 的读取侧延伸）
MAX_ROWS = 500
MAX_COLUMNS = 64
MAX_CELL_RECORDS = 5000


def _normalize(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


class XlsxAdapter(ArtifactAdapter):
    artifact_type = "xlsx"
    supported_extensions = (".xlsx",)
    parser_library = "openpyxl"

    def read(self, path: Path, artifact_id: str) -> Artifact:
        ensure_not_ole(path, self)
        # data_only=False 拿公式；data_only=True 拿 Excel 最近一次计算的缓存值
        wb_formulas = openpyxl.load_workbook(str(path), data_only=False)
        wb_values = openpyxl.load_workbook(str(path), data_only=True)

        blocks: list[ArtifactBlock] = []
        sheets_meta: list[dict] = []

        for sheet_name in wb_formulas.sheetnames:
            ws_formulas = wb_formulas[sheet_name]
            ws_values = wb_values[sheet_name]

            real_max_row = ws_formulas.max_row or 0
            real_max_col = ws_formulas.max_column or 0
            max_row = min(real_max_row, MAX_ROWS)
            max_col = min(real_max_col, MAX_COLUMNS)

            grid: list[list] = []
            cell_records: list[dict] = []
            cells_truncated = False
            for row in range(1, max_row + 1):
                grid_row: list = []
                for col in range(1, max_col + 1):
                    raw = ws_formulas.cell(row=row, column=col).value
                    formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                    if formula is not None:
                        value = _normalize(ws_values.cell(row=row, column=col).value)
                    else:
                        value = _normalize(raw)
                    grid_row.append("" if value is None else value)
                    if value is None and formula is None:
                        continue
                    if len(cell_records) >= MAX_CELL_RECORDS:
                        cells_truncated = True
                        continue
                    cell_records.append(
                        {
                            "address": f"{get_column_letter(col)}{row}",
                            "row": row,
                            "column": col,
                            "value": value,
                            "formula": formula,
                        }
                    )
                grid.append(grid_row)

            merged = [str(rng) for rng in ws_formulas.merged_cells.ranges]
            dimensions = ws_formulas.dimensions or "A1:A1"
            blocks.append(
                ArtifactBlock(
                    "",
                    "table",
                    None,
                    {"sheet": sheet_name, "range": dimensions},
                    {
                        "rows": max_row,
                        "columns": max_col,
                        "cells": grid,
                        "cell_records": cell_records,
                        "merged_cells": merged,
                        "rows_truncated": real_max_row > max_row,
                        "columns_truncated": real_max_col > max_col,
                        "cells_truncated": cells_truncated,
                    },
                )
            )
            sheets_meta.append(
                {
                    "name": sheet_name,
                    "dimensions": dimensions,
                    "max_row": real_max_row,
                    "max_column": real_max_col,
                    "merged_count": len(merged),
                    "non_empty_cells": len(cell_records),
                }
            )

        metadata = {
            "title": wb_formulas.properties.title or None,
            "sheet_count": len(sheets_meta),
            "sheets": sheets_meta,
            "active_sheet": (
                wb_formulas.active.title if wb_formulas.active is not None else None
            ),
        }

        content_lines: list[str] = []
        for block, sheet_meta in zip(blocks, sheets_meta):
            content_lines.append(f"[Sheet {sheet_meta['name']}] {sheet_meta['dimensions']}")
            for row in block.metadata["cells"][:50]:
                line = "\t".join("" if v is None else str(v) for v in row).rstrip("\t")
                if line:
                    content_lines.append(line)
        content = "\n".join(content_lines)

        assign_block_ids(blocks)
        return self.build_artifact(
            path,
            artifact_id,
            blocks=blocks,
            metadata=metadata,
            content=content,
            file_stat=path.stat(),
        )
